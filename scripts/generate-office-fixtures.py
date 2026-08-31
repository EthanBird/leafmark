#!/usr/bin/env python3
"""Generate real Word / PowerPoint fixtures for visual testing."""

from __future__ import annotations

from pathlib import Path
from xml.sax.saxutils import escape
from zipfile import ZIP_DEFLATED, ZipFile

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor as PptRgb
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt as PptPt

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "public" / "office-fixtures"
ASSETS = OUT / "_assets"
FONT = "/usr/share/fonts/truetype/wqy/wqy-microhei.ttc"


def font(size: int) -> ImageFont.FreeTypeFont:
    return ImageFont.truetype(FONT, size, index=0)


def paint_cover(path: Path) -> None:
    image = Image.new("RGB", (1280, 720), "#0f766e")
    draw = ImageDraw.Draw(image)
    for y in range(720):
        mix = int(15 + (y / 720) * 70)
        draw.line((0, y, 1280, y), fill=(15, 80 + mix // 4, 90 + mix // 3))
    draw.ellipse((820, -80, 1480, 580), fill="#14b8a6")
    draw.ellipse((70, 420, 420, 860), fill="#042f2e")
    image.save(path, "PNG", optimize=True)


def paint_photo(path: Path) -> None:
    image = Image.new("RGB", (640, 420), "#fed7aa")
    draw = ImageDraw.Draw(image)
    draw.rectangle((0, 260, 640, 420), fill="#7c2d12")
    draw.polygon([(0, 260), (180, 90), (340, 260)], fill="#fb923c")
    draw.polygon([(280, 260), (470, 40), (640, 260)], fill="#c2410c")
    draw.ellipse((480, 36, 600, 156), fill="#fef3c7")
    draw.text((24, 24), "插图", font=font(28), fill="#7c2d12")
    image.save(path, "PNG", optimize=True)


def paint_icon(path: Path) -> None:
    image = Image.new("RGB", (240, 240), "#1d4ed8")
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((28, 28, 212, 212), radius=36, fill="#93c5fd")
    draw.text((58, 88), "LM", font=font(64), fill="#1e3a8a")
    image.save(path, "PNG", optimize=True)


def add_hyperlink(paragraph, text: str, url: str) -> None:
    part = paragraph.part
    rel = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.append(color)
    rpr.append(underline)
    text_node = OxmlElement("w:t")
    text_node.set(qn("xml:space"), "preserve")
    text_node.text = text
    run.append(rpr)
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def float_picture(run, align: str = "right") -> None:
    drawing = run._element.find(qn("w:drawing"))
    inline = drawing.find(qn("wp:inline"))
    extent = inline.find(qn("wp:extent"))
    doc_pr = inline.find(qn("wp:docPr"))
    graphic = inline.find(qn("a:graphic"))
    anchor = OxmlElement("wp:anchor")
    for key, value in {
        "behindDoc": "0",
        "distT": "0",
        "distB": "0",
        "distL": "114300",
        "distR": "114300",
        "simplePos": "0",
        "relativeHeight": "251658240",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(key, value)
    simple = OxmlElement("wp:simplePos")
    simple.set("x", "0")
    simple.set("y", "0")
    pos_h = OxmlElement("wp:positionH")
    pos_h.set("relativeFrom", "column")
    align_el = OxmlElement("wp:align")
    align_el.text = align
    pos_h.append(align_el)
    pos_v = OxmlElement("wp:positionV")
    pos_v.set("relativeFrom", "paragraph")
    offset = OxmlElement("wp:posOffset")
    offset.text = "0"
    pos_v.append(offset)
    wrap = OxmlElement("wp:wrapSquare")
    wrap.set("wrapText", "bothSides")
    effect = OxmlElement("wp:effectExtent")
    for key in ("l", "t", "r", "b"):
        effect.set(key, "0")
    for child in (simple, pos_h, pos_v, extent, effect, wrap, doc_pr, graphic):
        if child is not None:
            anchor.append(child)
    drawing.remove(inline)
    drawing.append(anchor)


def write_docx(cover: Path, photo: Path, icon: Path, dest: Path) -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_heading("一叶办公文档视觉样张", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    intro = doc.add_paragraph()
    intro.add_run("这是一份用 python-docx 写出的真实 .docx，用来检查 LeafMark 能否显示图片、")
    add_hyperlink(intro, "打开官网链接", "https://github.com/EthanBird/leafmark")
    intro.add_run("、合并表格，以及右侧环绕的浮动图。正文应保持可读，不能变成标签源代码。")

    float_para = doc.add_paragraph()
    float_run = float_para.add_run()
    float_run.add_picture(str(icon), width=Cm(3.2))
    float_picture(float_run, "right")
    float_para.add_run(
        "右侧这张蓝色图标是浮动图（wp:anchor + wrapSquare）。文字应绕开它，而不是把它挤到下一页或显示成破图。"
        "如果编辑器只能处理 wp:inline，这里会看不见图或把图叠在字上。"
    )

    doc.add_heading("产品要点", level=2)
    for item in ("本地打开 OOXML，不上传", "未改动的 XML 原样写回 Word / WPS", "首屏只渲染可见内容"):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("检查顺序", level=2)
    for item in ("先确认浮动图只环绕本段，不漏到标题和表格", "再确认超链接是真实 URL，不是 rId", "最后看合并表格的跨行跨列"):
        doc.add_paragraph(item, style="List Number")

    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(photo), width=Cm(11.4))
    caption = doc.add_paragraph("图：内嵌插图（wp:inline），应完整显示山形色块与「插图」二字。")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in caption.runs:
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x5F, 0x63, 0x68)

    doc.add_heading("合并单元格", level=2)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    table.cell(0, 0).text = "模块"
    table.cell(0, 1).merge(table.cell(0, 2)).text = "说明（跨两列）"
    table.cell(1, 0).merge(table.cell(2, 0)).text = "编辑器\n（跨两行）"
    table.cell(1, 1).text = "Word"
    table.cell(1, 2).text = "图片 / 超链接 / 表格"
    table.cell(2, 1).text = "PPT"
    table.cell(2, 2).text = "背景图 / 表格 / 版式"

    note = doc.add_paragraph()
    note.add_run("页眉页脚也会写进文件，用于确认页眉带是否出现。").italic = True
    section.header.paragraphs[0].text = "LeafMark 视觉样张 · Word"
    section.footer.paragraphs[0].text = "只读检查用，不执行宏"

    dest.parent.mkdir(parents=True, exist_ok=True)
    doc.save(dest)


def set_run(paragraph, text: str, size: int, bold: bool = False, color: str = "202124") -> None:
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = PptRgb.from_string(color)
    paragraph.alignment = PP_ALIGN.LEFT


def write_pptx(cover: Path, photo: Path, dest: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    cover_slide = prs.slides.add_slide(blank)
    cover_slide.shapes.add_picture(str(cover), Emu(0), Emu(0), prs.slide_width, prs.slide_height)
    title_box = cover_slide.shapes.add_textbox(Inches(0.7), Inches(2.05), Inches(9.6), Inches(1.2))
    set_run(title_box.text_frame.paragraphs[0], "一叶演示文稿样张", 40, True, "F8FAFC")
    sub = cover_slide.shapes.add_textbox(Inches(0.7), Inches(3.35), Inches(9.6), Inches(0.7))
    set_run(sub.text_frame.paragraphs[0], "LeafMark 视觉样张 · 全幅背景图 + 标题文本框", 20, False, "CCFBF1")
    note = cover_slide.shapes.add_textbox(Inches(0.7), Inches(4.15), Inches(10.2), Inches(1.0))
    set_run(note.text_frame.paragraphs[0], "若背景丢失，这一页会变成白底。标题必须来自文本框，不能只印在 PNG 里。", 16, False, "99F6E4")

    content = prs.slides.add_slide(blank)
    fill = content.background.fill
    fill.solid()
    fill.fore_color.rgb = PptRgb.from_string("F8FAFC")
    heading = content.shapes.add_textbox(Inches(0.6), Inches(0.28), Inches(12), Inches(0.7))
    set_run(heading.text_frame.paragraphs[0], "图片与表格应同时可见", 28, True)
    content.shapes.add_picture(str(photo), Inches(0.6), Inches(1.2), Inches(5.6), Inches(3.7))
    table_shape = content.shapes.add_table(4, 3, Inches(6.6), Inches(1.25), Inches(6.0), Inches(4.0)).table
    table_shape.cell(0, 0).merge(table_shape.cell(0, 2))
    table_shape.cell(0, 0).text = "对照表（跨三列）"
    for col, value in enumerate(("项", "Word", "PPT")):
        table_shape.cell(1, col).text = value
    for col, value in enumerate(("图片", "内嵌 + 浮动", "全幅 + 插图")):
        table_shape.cell(2, col).text = value
    for col, value in enumerate(("表格", "合并单元格", "本页右侧")):
        table_shape.cell(3, col).text = value

    title_layout = prs.slide_layouts[0]
    layout_slide = prs.slides.add_slide(title_layout)
    layout_slide.shapes.title.text = "版式占位符标题"
    if len(layout_slide.placeholders) > 1:
        layout_slide.placeholders[1].text = "副标题：这一页来自 Title 版式。占位符几何应来自 slideLayout，而不是全部挤在左上角。"

    dest.parent.mkdir(parents=True, exist_ok=True)
    prs.save(dest)


def write_xlsx(dest: Path) -> None:
    strings = ["一叶表格视觉样张", "产品", "数量", "单价", "金额", "编辑器", "演示", "合计"]
    sst = "".join(f"<si><t>{escape(item)}</t></si>" for item in strings)
    shared = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        f'<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" count="{len(strings)}" uniqueCount="{len(strings)}">{sst}</sst>'
    )
    sheet = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <dimension ref="A1:D6"/>
  <sheetViews>
    <sheetView workbookViewId="0">
      <pane xSplit="0" ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>
    </sheetView>
  </sheetViews>
  <sheetFormatPr defaultRowHeight="18"/>
  <cols>
    <col min="1" max="1" width="18" customWidth="1"/>
    <col min="2" max="2" width="12" customWidth="1"/>
    <col min="3" max="3" width="12" customWidth="1"/>
    <col min="4" max="4" width="14" customWidth="1"/>
  </cols>
  <sheetData>
    <row r="1"><c r="A1" t="s"><v>0</v></c></row>
    <row r="2">
      <c r="A2" t="s"><v>1</v></c>
      <c r="B2" t="s"><v>2</v></c>
      <c r="C2" t="s"><v>3</v></c>
      <c r="D2" t="s"><v>4</v></c>
    </row>
    <row r="3">
      <c r="A3" t="s"><v>5</v></c>
      <c r="B3"><v>2</v></c>
      <c r="C3"><v>40</v></c>
      <c r="D3"><f>B3*C3</f><v>80</v></c>
    </row>
    <row r="4">
      <c r="A4" t="s"><v>6</v></c>
      <c r="B4"><v>3</v></c>
      <c r="C4"><v>25</v></c>
      <c r="D4"><f>B4*C4</f><v>75</v></c>
    </row>
    <row r="5">
      <c r="A5" t="s"><v>5</v></c>
      <c r="B5"><v>1</v></c>
      <c r="C5"><v>60</v></c>
      <c r="D5"><f>B5*C5</f><v>60</v></c>
    </row>
    <row r="6">
      <c r="A6" t="s"><v>7</v></c>
      <c r="D6"><f>SUM(D3:D5)</f><v>215</v></c>
    </row>
  </sheetData>
  <mergeCells count="1"><mergeCell ref="A1:D1"/></mergeCells>
</worksheet>
"""
    workbook = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="数据" sheetId="1" r:id="rId1"/></sheets>
</workbook>
"""
    rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>
"""
    workbook_rels = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/sharedStrings" Target="sharedStrings.xml"/>
</Relationships>
"""
    types = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/sharedStrings.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/>
</Types>
"""
    dest.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(dest, "w", ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", types)
        zf.writestr("_rels/.rels", rels)
        zf.writestr("xl/workbook.xml", workbook)
        zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
        zf.writestr("xl/worksheets/sheet1.xml", sheet)
        zf.writestr("xl/sharedStrings.xml", shared)


def main() -> None:
    ASSETS.mkdir(parents=True, exist_ok=True)
    cover = ASSETS / "cover.png"
    photo = ASSETS / "photo.png"
    icon = ASSETS / "icon.png"
    paint_cover(cover)
    paint_photo(photo)
    paint_icon(icon)
    write_docx(cover, photo, icon, OUT / "leafmark-sample.docx")
    write_pptx(cover, photo, OUT / "leafmark-sample.pptx")
    write_xlsx(OUT / "leafmark-sample.xlsx")
    print(f"wrote {OUT / 'leafmark-sample.docx'} ({(OUT / 'leafmark-sample.docx').stat().st_size} bytes)")
    print(f"wrote {OUT / 'leafmark-sample.pptx'} ({(OUT / 'leafmark-sample.pptx').stat().st_size} bytes)")
    print(f"wrote {OUT / 'leafmark-sample.xlsx'} ({(OUT / 'leafmark-sample.xlsx').stat().st_size} bytes)")


if __name__ == "__main__":
    main()
